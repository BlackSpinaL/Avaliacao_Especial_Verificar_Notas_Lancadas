import streamlit as st
import pandas as pd
import camelot
import re
import os
import hashlib
import tempfile
import unicodedata
import concurrent.futures
from io import BytesIO

st.set_page_config(page_title="Verificador AE x Diários", layout="wide")
st.title("🔍 Verificador de Lançamento de Notas — Avaliação Especial")

st.markdown("""
**Fluxo:** 1️⃣ envie a planilha → 2️⃣ veja quais diários são necessários →
3️⃣ **escolha uma turma** → 4️⃣ faça upload dos diários dela → 5️⃣ rode a verificação.
""")

COLUNAS_NOTAS = ["AP1/AV1", "AP2/AV2", "TE", "AE", "ND", "TOTAL PARCIAL", "FINAL"]

ALIASES_NOTAS = {
    "AP1/AV1":       ["AP1/AV1", "AV1/AP1", "AP1AV1", "AV1AP1"],
    "AP2/AV2":       ["AP2/AV2", "AV2/AP2", "AP2AV2", "AV2AP2",
                      "AP2/AS", "AP2AS", "AS/AP2"],
    "TE":            ["TE"],
    "AE":            ["AE"],
    "ND":            ["ND"],
    "TOTAL PARCIAL": ["TOTALPARCIAL"],
    "FINAL":         ["FINAL"],
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def normalizar(s):
    if s is None:
        return ""
    s = str(s).upper().strip()
    s = unicodedata.normalize("NFKD", s).encode("ASCII", "ignore").decode("ASCII")
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def normalizar_nome(s):
    """Normaliza nome para comparação: sem acento, sem espaço extra."""
    return normalizar(s)


def normalizar_celula(celula):
    if celula is None:
        return ""
    s = str(celula).upper()
    s = unicodedata.normalize("NFKD", s).encode("ASCII", "ignore").decode("ASCII")
    s = re.sub(r"[^A-Z0-9]", "", s)
    return s


def normalizar_codigo(v):
    if pd.isna(v):
        return ""
    try:
        return str(int(float(v)))
    except (ValueError, TypeError):
        return re.sub(r"\D", "", str(v))


def valor_em_branco(valor):
    if valor is None:
        return True
    s = str(valor).strip()
    if s in ("", "-", "--", "nan", "NaN", "None"):
        return True
    s_norm = s.replace(",", ".").replace(" ", "")
    try:
        return float(s_norm) == 0.0
    except ValueError:
        return False


def contar_colunas_notas(df_bruto):
    texto = " ".join(str(v) for v in df_bruto.values.flatten()).upper()
    texto_norm = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    texto_norm = re.sub(r"[^A-Z]", "", texto_norm)
    if "LABORATORIO" in texto_norm:
        return 9
    if "RECUPERACAO" in texto_norm or "RECUPERAC" in texto_norm:
        return 8
    return 7


def contar_valores_nao_zero(df_bruto, linhas_alunos, idx):
    n_cols = df_bruto.shape[1]
    if idx is None or idx < 0 or idx >= n_cols:
        return -1
    total = 0
    for row in linhas_alunos:
        v = str(row.iloc[idx]).strip().replace(",", ".")
        try:
            if float(v) != 0.0:
                total += 1
        except (ValueError, TypeError):
            pass
    return total


def descobrir_ae_idx(df_bruto, linhas_alunos, first_note_idx):
    n_cols = df_bruto.shape[1]
    idx_posicional = first_note_idx + 3

    idx_header = None
    for r in range(min(20, len(df_bruto))):
        for i in range(n_cols):
            if normalizar_celula(df_bruto.iloc[r, i]) == "AE":
                idx_header = i
                break
        if idx_header is not None:
            break

    candidatos = []
    if 0 <= idx_posicional < n_cols:
        candidatos.append(("posicional", idx_posicional))
    if idx_header is not None and idx_header != idx_posicional and 0 <= idx_header < n_cols:
        candidatos.append(("cabecalho", idx_header))

    if not candidatos:
        return idx_posicional
    if len(candidatos) == 1:
        return candidatos[0][1]

    candidatos_com_score = [
        (label, idx, contar_valores_nao_zero(df_bruto, linhas_alunos, idx))
        for label, idx in candidatos
    ]
    candidatos_com_score.sort(key=lambda x: x[2], reverse=True)
    return candidatos_com_score[0][1]


def extrair_tabela_notas(df_bruto):
    linhas_alunos = []
    for _, row in df_bruto.iterrows():
        primeira = str(row.iloc[0]).strip()
        if re.match(r"^\d{1,3}$", primeira):
            linhas_alunos.append(row)

    if not linhas_alunos:
        return None

    n_cols = df_bruto.shape[1]
    if n_cols < 4:
        return None

    n_notas = contar_colunas_notas(df_bruto)
    if n_notas > n_cols - 2:
        n_notas = 7
    first_note_idx = n_cols - n_notas
    ae_idx = descobrir_ae_idx(df_bruto, linhas_alunos, first_note_idx)

    def safe_get(row, idx):
        if idx is None or idx < 0 or idx >= n_cols:
            return ""
        return row.iloc[idx]

    registros = []
    for row in linhas_alunos:
        # Tenta matrícula em iloc[1]; se vier vazio, tenta iloc[2] e usa o nome de iloc[3]
        matricula_raw = str(row.iloc[1]).strip()
        matricula = re.sub(r"\D", "", matricula_raw)
        nome = str(row.iloc[2]).strip()

        if not matricula:
            # fallback: talvez o Camelot tenha deslocado colunas
            matricula = re.sub(r"\D", "", str(row.iloc[2]))
            nome = str(row.iloc[3]).strip() if len(row) > 3 else ""

        if not matricula:
            continue

        reg = {"MATRICULA": matricula, "NOME": nome}
        reg["AP1/AV1"]       = safe_get(row, first_note_idx + 0)
        reg["AP2/AV2"]       = safe_get(row, first_note_idx + 1)
        reg["TE"]            = safe_get(row, first_note_idx + 2)
        reg["AE"]            = safe_get(row, ae_idx)
        reg["ND"]            = safe_get(row, first_note_idx + 4)
        reg["TOTAL PARCIAL"] = safe_get(row, n_cols - 2)
        reg["FINAL"]         = safe_get(row, n_cols - 1)
        registros.append(reg)

    return pd.DataFrame(registros) if registros else None


def eh_tabela_de_notas(df_bruto):
    texto = " ".join(str(v) for v in df_bruto.values.flatten()).upper()
    texto_norm = re.sub(r"\s+", "", texto)
    marcadores = ["AP1/AV1", "AV1/AP1", "AP1AV1", "AV1AP1",
                  "AP2/AS", "AP2/AV2", "TOTALPARCIAL"]
    return sum(1 for m in marcadores if m in texto_norm) >= 2


def _tentar_flavor(tmp_path: str, **kwargs):
    try:
        tabs = camelot.read_pdf(tmp_path, pages="all", **kwargs)
    except Exception as e:
        return {"erro": str(e), "notas": {}, "tabelas": [], "n_notas_max": 0}

    notas = {}
    n_notas_max = 0
    tabelas_info = []
    for i, t in enumerate(tabs):
        df = t.df
        n_cols = df.shape[1]
        n_notas = contar_colunas_notas(df)
        first_note_idx = n_cols - n_notas

        idx_header = None
        for r in range(min(20, len(df))):
            for j in range(n_cols):
                if normalizar_celula(df.iloc[r, j]) == "AE":
                    idx_header = j
                    break
            if idx_header is not None:
                break

        eh_tab = eh_tabela_de_notas(df)
        tabelas_info.append({
            "indice": i,
            "shape": df.shape,
            "n_notas_detectado": n_notas,
            "first_note_idx": first_note_idx,
            "AE_idx_posicional": first_note_idx + 3,
            "AE_idx_cabecalho": idx_header,
            "eh_tabela_notas": eh_tab,
        })
        if not eh_tab:
            continue
        df_alunos = extrair_tabela_notas(df)
        if df_alunos is None:
            continue
        for _, row in df_alunos.iterrows():
            notas[row["MATRICULA"]] = row.to_dict()
        n_notas_max = max(n_notas_max, n_notas)

    return {"erro": None, "notas": notas, "tabelas": tabelas_info, "n_notas_max": n_notas_max}


@st.cache_data(show_spinner=False)
def processar_pdf(pdf_bytes: bytes) -> dict:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(pdf_bytes)
        tmp_path = tmp.name
    try:
        configs = [
            ("lattice_40", {"flavor": "lattice", "line_scale": 40}),
            ("stream",     {"flavor": "stream",  "strip_text": "\n"}),
        ]
        debug = {}
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as ex:
            futuros = {
                ex.submit(_tentar_flavor, tmp_path, **kwargs): label
                for label, kwargs in configs
            }
            for fut in concurrent.futures.as_completed(futuros):
                label = futuros[fut]
                try:
                    debug[label] = fut.result()
                except Exception as e:
                    debug[label] = {"erro": str(e), "notas": {}, "tabelas": [], "n_notas_max": 0}

        l40 = debug.get("lattice_40", {})
        if l40.get("notas") and len(l40["notas"]) >= 20:
            return {"notas": l40["notas"], "debug": debug, "vencedor": "lattice_40"}

        opcoes = [
            (label, len(d.get("notas", {})), d.get("n_notas_max", 0), d.get("notas", {}))
            for label, d in debug.items() if d.get("notas")
        ]
        if not opcoes:
            return {"notas": {}, "debug": debug, "vencedor": None}
        opcoes.sort(key=lambda x: (x[1], x[2]), reverse=True)
        vencedor = opcoes[0][0]
        return {"notas": opcoes[0][3], "debug": debug, "vencedor": vencedor}
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass


def extrair_notas_pdf(pdf_bytes: bytes) -> dict:
    return processar_pdf(pdf_bytes).get("notas", {})


def depurar_pdf(pdf_bytes: bytes) -> dict:
    return processar_pdf(pdf_bytes).get("debug", {})


def encontrar_aluno(notas: dict, matricula: str, nome: str):
    """
    Procura o aluno em `notas` por 3 estratégias:
      1) match exato de matrícula
      2) match flexível (lstrip de zeros) de matrícula
      3) match por nome normalizado
    Devolve o dict de notas ou None.
    """
    # 1) exato
    if matricula in notas:
        return notas[matricula]

    # 2) lstrip de zeros
    mat_limpa = matricula.lstrip("0")
    for m, n in notas.items():
        if m.lstrip("0") == mat_limpa:
            return n

    # 3) por nome normalizado
    nome_norm = normalizar_nome(nome)
    if not nome_norm:
        return None
    for m, n in notas.items():
        nome_pdf = normalizar_nome(n.get("NOME", ""))
        if nome_pdf and nome_pdf == nome_norm:
            return n
    return None


def aplicar_formatacao_excel(writer, sheet_name, coluna_status):
    from openpyxl.styles import PatternFill, Font
    ws = writer.sheets[sheet_name]
    header = [c.value for c in ws[1]]
    if coluna_status not in header:
        return
    col_idx = header.index(coluna_status) + 1
    verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    laranja = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    cinza = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    verde_f = Font(color="006100")
    vermelho_f = Font(color="9C0006")
    laranja_f = Font(color="9C6500")
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        v = str(cell.value).strip() if cell.value is not None else ""
        if v == "Sim":
            cell.fill, cell.font = verde, verde_f
        elif v == "Não":
            cell.fill, cell.font = vermelho, vermelho_f
        elif v in ("Diário não enviado", "Aluno não encontrado"):
            cell.fill, cell.font = laranja, laranja_f
        elif v == "":
            cell.fill = cinza
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 60)


def gerar_planilha_modelo() -> bytes:
    modelo = pd.DataFrame(columns=[
        "Matrícula", "Nome do Aluno", "Turma", "Etapa",
        "1ª Disciplina", "2ª Disciplina", "3ª Disciplina", "4ª Disciplina",
    ])
    modelo.loc[0] = ["000000", "AAAAA AAAAA AAAAA", "12101", "2",
                     "019 - CIENCIAS", "009 - GEOGRAFIA",
                     "276 - LIN.PORTUGUESA 2", ""]
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        modelo.to_excel(writer, sheet_name="Lista de Solicitação", index=False)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# ETAPA 1 — Upload da planilha
# ---------------------------------------------------------------------------
st.header("1️⃣ Lista de solicitações")
st.warning(
    "⚠️ **ATENÇÃO** — A planilha deve conter **SOMENTE** as colunas do modelo abaixo.\n\n"
    "Colunas extras (filtros, observações, notas antigas, 'Disciplina anterior', etc.) "
    "podem atrapalhar a varredura."
)

col_img, col_txt = st.columns([3, 2])
with col_img:
    caminho_imagem = "assets/template_solicitacoes.png"
    if os.path.exists(caminho_imagem):
        st.image(caminho_imagem, caption="Modelo correto da planilha",
                 use_container_width=True)
    else:
        st.markdown("**Modelo correto (exemplo):**")
        df_exemplo = pd.DataFrame({
            "Matrícula": ["000000", "111111"],
            "Nome do Aluno": ["AAAAA AAAAA AAAAA", "BBBBB BBBBB BBBBB"],
            "Turma": ["12101", "12102"],
            "Etapa": ["2", "2"],
            "1ª Disciplina": ["019 - CIENCIAS", "017 - MATEMATICA"],
            "2ª Disciplina": ["009 - GEOGRAFIA", "009 - GEOGRAFIA"],
            "3ª Disciplina": ["276 - LIN.PORTUGUESA 2", ""],
            "4ª Disciplina": ["", ""],
        })
        st.dataframe(df_exemplo, hide_index=True, use_container_width=True)

with col_txt:
    st.markdown(
        """
        **Colunas obrigatórias:**
        - `Matrícula`
        - `Nome do Aluno`
        - `Turma`
        - `Etapa`
        - Pelo menos uma coluna de disciplina (`1ª Disciplina`, `2ª Disciplina`, ...)

        **Formato da disciplina:** `código - NOME` (ex.: `017 - MATEMATICA`)

        **Remova da planilha** qualquer coluna que não esteja no modelo —
        inclusive filtros e colunas em branco no final.
        """
    )

st.download_button(
    label="📄 Baixar planilha-modelo (.xlsx)",
    data=gerar_planilha_modelo(),
    file_name="modelo_solicitacoes_avaliacao_especial.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)

st.markdown("---")

uploaded_excel = st.file_uploader(
    "Envie o arquivo Excel com as solicitações", type=["xlsx"], key="excel"
)
if not uploaded_excel:
    st.info("⬆️ Envie a planilha para começar.")
    st.stop()

excel_hash = hashlib.md5(uploaded_excel.getvalue()).hexdigest()[:8]
df_excel = pd.read_excel(uploaded_excel)

colunas_obrigatorias = ["Matrícula", "Nome do Aluno", "Turma", "Etapa"]
colunas_presentes = [str(c).strip() for c in df_excel.columns]
faltando_obrig = [c for c in colunas_obrigatorias if c not in colunas_presentes]

padrao_disc = re.compile(r"^\s*[1-4]ª\s*Disciplina\s*$", re.IGNORECASE)
disc_cols = [c for c in df_excel.columns if padrao_disc.match(str(c))]
colunas_disc_extra = [
    c for c in df_excel.columns
    if "Disciplina" in str(c) and not padrao_disc.match(str(c))
]
colunas_permitidas = set(colunas_obrigatorias) | set(disc_cols) | set(colunas_disc_extra)
colunas_extras = [c for c in df_excel.columns if c not in colunas_permitidas]

if faltando_obrig:
    st.error(f"❌ Faltam colunas obrigatórias: **{', '.join(faltando_obrig)}**.")
    st.stop()
if not disc_cols:
    st.error("❌ Nenhuma coluna de disciplina encontrada.")
    st.stop()
if colunas_disc_extra:
    st.error(f"❌ Colunas fora do padrão: **{', '.join(map(str, colunas_disc_extra))}**.")
    st.stop()
if colunas_extras:
    st.warning(f"⚠️ Colunas extras ignoradas: **{', '.join(map(str, colunas_extras))}**.")

st.success(f"✅ Planilha validada — disciplinas: **{', '.join(map(str, disc_cols))}**")

solicitacoes = []
for _, row in df_excel.iterrows():
    matricula = normalizar_codigo(row.get("Matrícula"))
    nome = str(row.get("Nome do Aluno", "")).strip()
    turma = normalizar_codigo(row.get("Turma"))
    etapa = str(row.get("Etapa", "")).strip()
    for i, col in enumerate(disc_cols, start=1):
        val = row.get(col)
        if pd.isna(val) or str(val).strip() == "":
            continue
        disc_str = str(val).strip()
        if " - " in disc_str:
            partes = disc_str.split(" - ", 1)
            disc_code = partes[0].strip()
            disc_name = partes[1].strip()
        else:
            disc_code = ""
            disc_name = disc_str
        solicitacoes.append({
            "matricula": matricula, "nome": nome, "turma": turma,
            "etapa": etapa, "posicao": i,
            "disciplina_raw": disc_str,
            "disciplina_code": disc_code,
            "disciplina_name": disc_name,
        })

if not solicitacoes:
    st.error("Nenhuma solicitação encontrada.")
    st.stop()

# ---------------------------------------------------------------------------
# ETAPA 2 — Diários necessários
# ---------------------------------------------------------------------------
st.header("2️⃣ Diários que devem ser enviados")

necessarios = {}
for sol in solicitacoes:
    key = (sol["turma"], sol["disciplina_code"])
    if key not in necessarios:
        necessarios[key] = {"nome": sol["disciplina_name"], "qtd": 0}
    necessarios[key]["qtd"] += 1

linhas_resumo = []
for (turma, code), info in sorted(necessarios.items()):
    linhas_resumo.append({
        "Turma": turma, "Disciplina": info["nome"],
        "Código": code or "—", "Solicitações": info["qtd"],
    })
df_resumo = pd.DataFrame(linhas_resumo)
st.dataframe(df_resumo, use_container_width=True, hide_index=True)
st.caption(f"📊 {len(solicitacoes)} solicitação(ões) em {len(necessarios)} diário(s).")

# ---------------------------------------------------------------------------
# ETAPA 3 — Seleção de UMA turma + uploads
# ---------------------------------------------------------------------------
st.header("3️⃣ Escolha a turma que você vai verificar")
st.markdown(
    "Selecione **uma turma por vez**. O app processa apenas os diários "
    "dessa turma — fica mais rápido e facilita isolar eventuais problemas."
)

turmas = sorted(set(s["turma"] for s in solicitacoes))

select_key = f"turma_sel_{excel_hash}"
if select_key not in st.session_state:
    st.session_state[select_key] = "—"

opcoes = ["—"] + turmas
turma_escolhida = st.selectbox(
    "Turma a processar",
    opcoes,
    key=select_key,
    help="Escolha a turma cujos diários você vai enviar agora.",
)

if turma_escolhida == "—":
    st.info("☝️ Escolha uma turma no seletor acima para liberar os uploads.")
    st.stop()

st.success(f"🎯 Turma selecionada: **{turma_escolhida}**")

discs_turma = sorted(
    [(c, d) for (t, c), d in necessarios.items() if t == turma_escolhida],
    key=lambda x: x[1]["nome"],
)
qtd_sol_turma = sum(d["qtd"] for _, d in discs_turma)

st.markdown(
    f"### Diários da turma {turma_escolhida} — "
    f"{len(discs_turma)} diário(s) • {qtd_sol_turma} solicitação(ões)"
)

diarios_carregados = {}
cols = st.columns(min(len(discs_turma), 3)) if len(discs_turma) > 1 else [st]
for idx, (code, info) in enumerate(discs_turma):
    container = cols[idx % len(cols)]
    with container:
        st.markdown(
            f"**{info['nome']}**  \n"
            f"<small>{info['qtd']} solicitação(ões) — código {code or '—'}</small>",
            unsafe_allow_html=True,
        )
        up = st.file_uploader(
            f"Upload {info['nome']} (turma {turma_escolhida})",
            type=["pdf"],
            key=f"pdf_{excel_hash}_{turma_escolhida}_{code}",
            label_visibility="collapsed",
        )
        if up is not None:
            diarios_carregados[(turma_escolhida, code)] = {
                "arquivo": up, "nome": info["nome"],
            }

total_necessarios_sel = len(discs_turma)
faltando = total_necessarios_sel - len(diarios_carregados)

c1, c2, c3 = st.columns(3)
c1.metric("Diários esperados", total_necessarios_sel)
c2.metric("Diários carregados", len(diarios_carregados))
c3.metric("Faltando", max(0, faltando))

if faltando > 0:
    st.warning(f"⚠️ Faltam **{faltando}** diário(s).")
else:
    st.success("✅ Todos os diários da turma foram enviados!")

# ---------------------------------------------------------------------------
# ETAPA 4 — Verificação
# ---------------------------------------------------------------------------
st.header("4️⃣ Rodar verificação")

if st.button("▶️ Rodar verificação", type="primary"):

    solicitacoes_sel = [s for s in solicitacoes if s["turma"] == turma_escolhida]

    diarios_notas = {}
    total_diarios = len(diarios_carregados)
    if total_diarios > 0:
        progresso = st.progress(0, text="Preparando leitura dos PDFs...")
        for i, ((turma, code), info) in enumerate(diarios_carregados.items(), start=1):
            progresso.progress(
                i / total_diarios,
                text=f"Lendo {i}/{total_diarios} — {info['nome']}"
            )
            try:
                notas = extrair_notas_pdf(info["arquivo"].getvalue())
                diarios_notas[(turma, code)] = notas
            except Exception as e:
                st.warning(f"⚠️ Erro em {turma}/{info['nome']}: {e}")
        progresso.empty()

    # Estatísticas de match para o painel de depuração
    match_stats = {}  # (turma, code) -> {"exato": N, "flex": N, "nome": N, "falha": N}

    for sol in solicitacoes_sel:
        key = (sol["turma"], sol["disciplina_code"])
        if key not in diarios_notas:
            sol["status"] = "Diário não enviado"
            continue
        notas = diarios_notas[key]
        n_aluno = encontrar_aluno(notas, sol["matricula"], sol["nome"])
        if n_aluno is None:
            sol["status"] = "Aluno não encontrado"
            match_stats.setdefault(key, {"exato": 0, "flex": 0, "nome": 0, "falha": 0})
            match_stats[key]["falha"] += 1
        else:
            # Classifica o tipo de match (só para debug)
            if sol["matricula"] in notas:
                tipo = "exato"
            elif any(m.lstrip("0") == sol["matricula"].lstrip("0") for m in notas):
                tipo = "flex"
            else:
                tipo = "nome"
            match_stats.setdefault(key, {"exato": 0, "flex": 0, "nome": 0, "falha": 0})
            match_stats[key][tipo] += 1

            ae = n_aluno.get("AE", "")
            sol["status"] = "Não" if valor_em_branco(ae) else "Sim"

    linhas_long = []
    for sol in solicitacoes_sel:
        linhas_long.append({
            "Matrícula": sol["matricula"], "Nome do Aluno": sol["nome"],
            "Turma": sol["turma"], "Etapa": sol["etapa"],
            "Disciplina": sol["disciplina_raw"], "Lançou a Nota?": sol["status"],
        })
    df_long = (
        pd.DataFrame(linhas_long)
        .sort_values(by=["Nome do Aluno", "Disciplina"])
        .reset_index(drop=True)
    )

    alunos = {}
    for sol in solicitacoes_sel:
        key = (sol["turma"], sol["matricula"])
        if key not in alunos:
            alunos[key] = {
                "Matrícula": sol["matricula"], "Nome do Aluno": sol["nome"],
                "Turma": sol["turma"], "Etapa": sol["etapa"], "disciplinas": {},
            }
        alunos[key]["disciplinas"][sol["posicao"]] = sol

    linhas_wide = []
    for _, alu in alunos.items():
        linha = {
            "Matrícula": alu["Matrícula"], "Nome do Aluno": alu["Nome do Aluno"],
            "Turma": alu["Turma"], "Etapa": alu["Etapa"],
        }
        for i in range(1, 5):
            disc = alu["disciplinas"].get(i)
            if disc:
                linha[f"{i}ª Disciplina"] = disc["disciplina_raw"]
                linha[f"Lançou a Nota da {i}ª Disciplina? (Sim / Não)"] = disc["status"]
            else:
                linha[f"{i}ª Disciplina"] = ""
                linha[f"Lançou a Nota da {i}ª Disciplina? (Sim / Não)"] = ""
        linhas_wide.append(linha)

    df_wide = (
        pd.DataFrame(linhas_wide)
        .sort_values(by=["Nome do Aluno"])
        .reset_index(drop=True)
    )

    st.markdown("---")
    st.subheader(f"📈 Resumo — Turma {turma_escolhida}")
    total_sol = len(solicitacoes_sel)
    total_sim = sum(1 for s in solicitacoes_sel if s["status"] == "Sim")
    total_nao = sum(1 for s in solicitacoes_sel if s["status"] == "Não")
    total_sem_diario = sum(1 for s in solicitacoes_sel if s["status"] == "Diário não enviado")
    total_sem_aluno = sum(1 for s in solicitacoes_sel if s["status"] == "Aluno não encontrado")

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Total", total_sol)
    c2.metric("✅ Lançadas", total_sim)
    c3.metric("❌ Não lançadas", total_nao)
    c4.metric("📄 Sem diário", total_sem_diario)
    c5.metric("👤 Sem aluno", total_sem_aluno)

    st.markdown("---")
    tab1, tab2 = st.tabs(["📊 Por solicitação", "📋 Por aluno"])
    with tab1:
        st.dataframe(df_long, use_container_width=True, hide_index=True)
    with tab2:
        st.dataframe(df_wide, use_container_width=True, hide_index=True)

    with st.expander("🔎 Depuração da leitura dos PDFs (clique para abrir)"):
        st.caption("Cada PDF é processado uma única vez, com 2 configs em paralelo.")
        for (turma, code), info in diarios_carregados.items():
            st.markdown(f"### {info['nome']} (código {code})")
            notas_diario = diarios_notas.get((turma, code), {})
            st.write(f"**Alunos encontrados no PDF:** {len(notas_diario)}")

            # Estatísticas de match
            st_diario = match_stats.get((turma, code))
            if st_diario:
                st.write(
                    f"**Matches** — exato: {st_diario['exato']} • "
                    f"flex (lstrip 0): {st_diario['flex']} • "
                    f"por nome: {st_diario['nome']} • "
                    f"falhas: {st_diario['falha']}"
                )

            # Lista de matrículas extraídas (para conferência manual)
            if notas_diario:
                amostra = list(notas_diario.keys())
                st.write(f"**Matrículas extraídas ({len(amostra)}):** {amostra}")

            try:
                dbg = depurar_pdf(info["arquivo"].getvalue())
            except Exception as e:
                st.error(f"Erro ao depurar: {e}")
                continue
            for flavor, res in dbg.items():
                st.markdown(f"**Config `{flavor}`** — {res['erro'] or 'ok'}")
                if res["erro"]:
                    continue
                for t in res["tabelas"]:
                    st.write(
                        f"- Tabela #{t['indice']} — shape={t['shape']} — "
                        f"é tabela de notas? **{t['eh_tabela_notas']}**"
                    )
                    st.write(
                        f"  - colunas de notas detectadas: **{t['n_notas_detectado']}**"
                    )
                    st.write(
                        f"  - AE posicional: {t['AE_idx_posicional']} "
                        f"| AE cabeçalho: {t['AE_idx_cabecalho']}"
                    )

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_long.to_excel(writer, sheet_name=f"Turma {turma_escolhida}"[:31], index=False)
        aplicar_formatacao_excel(writer, f"Turma {turma_escolhida}"[:31], "Lançou a Nota?")
        sheet_wide = f"Turma {turma_escolhida} (aluno)"[:31]
        df_wide.to_excel(writer, sheet_name=sheet_wide, index=False)
        for i in range(1, 5):
            col_status = f"Lançou a Nota da {i}ª Disciplina? (Sim / Não)"
            aplicar_formatacao_excel(writer, sheet_wide, col_status)

    st.download_button(
        label=f"📥 Baixar resultado da turma {turma_escolhida} (Excel)",
        data=output.getvalue(),
        file_name=f"resultado_varredura_turma_{turma_escolhida}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
